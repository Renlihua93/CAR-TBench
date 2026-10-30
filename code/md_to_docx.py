# -*- coding: utf-8 -*-
"""
Minimal Markdown -> DOCX converter for the manuscript draft.

Handles headings, paragraphs, bullet lists and pipe tables. This is an
editable Word version, not a journal-formatted template; journal-specific
styling still has to be applied in Word.

Usage:
  python md_to_docx.py manuscript/JBI_submission_v1.md
"""
from __future__ import annotations

import argparse
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def add_runs(par, text: str):
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2])
            r.bold = True
        else:
            par.add_run(part)


def is_table_row(line: str) -> bool:
    return line.strip().startswith("|") and line.strip().endswith("|")


def cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("--out", default="")
    a = ap.parse_args()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(1.0)

    lines = open(a.src, encoding="utf-8").read().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(
                r"^\|[\s:\-|]+\|?$", lines[i + 1].strip()):
            header = cells(line)
            rows = []
            j = i + 2
            while j < len(lines) and is_table_row(lines[j]):
                rows.append(cells(lines[j]))
                j += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            for k, h in enumerate(header):
                cell = table.rows[0].cells[k]
                cell.text = ""
                add_runs(cell.paragraphs[0], h)
                for r in cell.paragraphs[0].runs:
                    r.bold = True
            for row in rows:
                cs = table.add_row().cells
                for k in range(len(header)):
                    cs[k].text = ""
                    add_runs(cs[k].paragraphs[0], row[k] if k < len(row) else "")
            doc.add_paragraph()
            i = j
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            doc.add_heading(m.group(2).strip(), level=min(level, 4))
            i += 1
            continue
        if line.strip().startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, line.strip()[2:])
            i += 1
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_runs(p, line.strip())
        i += 1

    out = a.out or os.path.splitext(a.src)[0] + ".docx"
    doc.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
